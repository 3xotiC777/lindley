import openpyxl
import pandas as pd
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config


def write_df_to_sheet(wb, sheet_name, df):
    ws = wb.create_sheet(title=sheet_name)
    clean = df.copy().fillna("")
    width_overrides = {
        "Control": 38, "Motivo": 70, "NOMBRE": 16, "Subcanal2": 14, "RUT.COM": 16,
        "ESTRATEGICA": 14, "LOC.COM": 12, "Objetivo": 12, "Titulares": 12,
        "Diferencia": 12, "Historico_%": 14, "Actual_%": 12, "Variacion_pp": 14,
        "Disponibles_SI": 15, "Obligatorios_SI": 16, "Referencia_30": 15, "Referencia_90": 15,
        "Rutas_ON_totales": 17, "Sin_candidato_OFF": 18, "Estado": 12,
        "Sin_candidato": 16,
        "Historico_ON": 14, "Objetivo_ON": 14, "Historico_OFF": 14, "Objetivo_OFF": 14,
    }
    for index, column in enumerate(clean.columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width_overrides.get(str(column), min(max(len(str(column)) + 2, 12), 28))
    header = []
    for column in clean.columns:
        cell = WriteOnlyCell(ws, value=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(cell)
    ws.append(header)
    ws.freeze_panes = "A2"
    for row in clean.itertuples(index=False, name=None):
        ws.append(list(row))


def export_selection_workbook(df_elegible, df_no_elegible, controls, output_path):
    """Export eligibility, selections and an auditable quota-control sheet."""
    print(f"Exporting main selection workbook to: {output_path}")
    wb = openpyxl.Workbook(write_only=True)
    write_df_to_sheet(wb, "BD", df_elegible)
    write_df_to_sheet(wb, "CONTROL CUOTAS", controls)

    off = df_elegible[df_elegible["_CANAL_SELECCION"].eq("OFF")]
    mix_rows = []
    for city, group in off[off[config.SELECTION_COL].eq("T")].groupby("NOMBRE"):
        total = len(group)
        mix_rows.append({
            "Ciudad": city,
            "Total_T": total,
            "EG": int(group["ESTRATEGICA"].astype(str).str.upper().str.startswith("EG").sum()),
            "GR": int(group["ESTRATEGICA"].astype(str).str.upper().str.startswith("GR").sum()),
            "ME": int(group["ESTRATEGICA"].astype(str).str.upper().str.startswith("ME").sum()),
            "PE": int(group["ESTRATEGICA"].astype(str).str.upper().str.startswith("PE").sum()),
        })
    write_df_to_sheet(wb, "VAR - MIX", pd.DataFrame(mix_rows))

    lima_controls = controls[controls["Control"].eq("Mínimo Lima por LOC.COM")]
    write_df_to_sheet(wb, "CDA LIMA", lima_controls)
    write_df_to_sheet(wb, "NO ELEGIBLES", df_no_elegible)
    wb.save(output_path)
    print("Main selection workbook saved successfully.")


def export_def_sup_workbook(df_elegible, output_path):
    print(f"Exporting ARCHIVO_DEF_SUP to: {output_path}")
    wb = openpyxl.Workbook(write_only=True)
    columns = ["CODIGO", "RUT.COM", "Ventas", config.SELECTION_COL]
    for canal in ("OFF", "ON"):
        export = df_elegible[df_elegible["_CANAL_SELECCION"].eq(canal)][columns].copy()
        export.columns = ["CODIGO", "RUT.COM", "Ventas", "seleccion"]
        write_df_to_sheet(wb, canal, export)
    all_rows = df_elegible[columns].copy()
    all_rows.columns = ["CODIGO", "RUT.COM", "Ventas", "seleccion"]
    write_df_to_sheet(wb, "Hoja4", all_rows)
    wb.save(output_path)
    print("ARCHIVO_DEF_SUP workbook saved successfully.")
